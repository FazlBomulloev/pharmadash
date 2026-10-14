import pytest
from openpyxl import Workbook

from backend.services.parsers.directions_parser import (
    DirectionsParseError,
    SHEET_NAME,
    parse_directions,
)


def make_file(tmp_path, sheet_title, rows, header=True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    if header:
        ws.append(["МНН", "ЛФ", "Направление", "Процент"])
    for r in rows:
        ws.append(list(r))
    path = tmp_path / "ref.xlsx"
    wb.save(path)
    return path


def test_parses_named_sheet(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("CLOBETASOL", "D-CRM", "Дерматология", 0.85),
        ("CLOBETASOL", "D-CRM", "Педиатрия", 0.15),
    ])
    rows, report = parse_directions(path)

    assert len(rows) == 2
    assert rows[0] == {
        "mnn": "CLOBETASOL", "lf": "D-CRM",
        "direction": "Дерматология", "share": 0.85,
    }
    assert report["sheet"] == SHEET_NAME
    assert report["rows_accepted"] == 2
    assert report["pairs_count"] == 1
    assert report["directions_count"] == 2
    assert report["forms_count"] == 1
    assert report["pairs_incomplete"] == 0


def test_finds_sheet_by_columns_when_name_differs(tmp_path):
    path = make_file(tmp_path, "Данные", [
        ("CLOBETASOL", "D-CRM", "Дерматология", 1.0),
    ])
    rows, report = parse_directions(path)

    assert len(rows) == 1
    assert report["sheet"] == "Данные"


def test_raises_when_no_suitable_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Прочее"
    ws.append(["Колонка", "Другая"])
    ws.append(["значение", 1])
    path = tmp_path / "bad.xlsx"
    wb.save(path)

    with pytest.raises(DirectionsParseError) as exc:
        parse_directions(path)

    assert "Длинный формат" in str(exc.value)


def test_percent_above_one_is_divided_by_hundred(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 35),
        ("A", "TC-TAB", "Педиатрия", 65),
    ])
    rows, report = parse_directions(path)

    assert [r["share"] for r in rows] == [0.35, 0.65]
    assert report["pairs_incomplete"] == 0


def test_skips_non_positive_and_empty_rows(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 1.0),
        ("A", "TC-TAB", "Педиатрия", 0),
        ("A", "TC-TAB", "ЛОР", None),
        (None, None, None, None),
        ("", "TC-TAB", "Терапия", 0.5),
        ("B", "TC-TAB", "", 0.5),
    ])
    rows, report = parse_directions(path)

    assert len(rows) == 1
    assert report["rows_skipped"] == 4


def test_normalises_case_and_whitespace(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        (" clobetasol ", " d-crm ", "  Дерматология  ", 1.0),
    ])
    rows, _ = parse_directions(path)

    assert rows[0]["mnn"] == "CLOBETASOL"
    assert rows[0]["lf"] == "D-CRM"
    assert rows[0]["direction"] == "Дерматология"


def test_duplicate_triple_keeps_last(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.4),
        ("A", "TC-TAB", "Терапия", 0.9),
    ])
    rows, report = parse_directions(path)

    assert len(rows) == 1
    assert rows[0]["share"] == 0.9
    assert report["rows_accepted"] == 1
    assert report["rows_duplicate"] == 1


def test_counts_pairs_with_incomplete_shares(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.5),
        ("A", "TC-TAB", "Педиатрия", 0.4),
        ("B", "TC-TAB", "Терапия", 1.0),
    ])
    _, report = parse_directions(path)

    assert report["pairs_incomplete"] == 1


def test_single_share_above_hundred_is_dropped_by_value_not_by_pair(tmp_path):
    # 150 -> 1.5: отбрасывается проверкой диапазона значения. Без неё строку
    # забрал бы фильтр пар, и pairs_over_100 стал бы 1.
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 150),
    ])
    rows, report = parse_directions(path)

    assert rows == []
    assert report["rows_skipped"] == 1
    assert report["pairs_over_100"] == 0


def test_pair_whose_shares_exceed_hundred_is_dropped(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.7),
        ("A", "TC-TAB", "Педиатрия", 0.6),
        ("B", "TC-TAB", "Терапия", 1.0),
    ])
    rows, report = parse_directions(path)

    assert [r["mnn"] for r in rows] == ["B"]
    assert report["pairs_over_100"] == 1
    assert report["pairs_count"] == 1
    assert report["rows_accepted"] == 1


def test_out_of_range_value_is_dropped_without_killing_its_pair(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 150),       # 1.5 после деления — вне диапазона
        ("A", "TC-TAB", "Педиатрия", 0.2),
    ])
    rows, report = parse_directions(path)

    # битая строка отброшена по значению, пара выжила с остальными долями
    assert [r["direction"] for r in rows] == ["Педиатрия"]
    assert report["rows_skipped"] == 1
    assert report["pairs_over_100"] == 0
    assert report["pairs_count"] == 1


def test_report_arithmetic_closes(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.6),
        ("A", "TC-TAB", "Педиатрия", 0.4),
        ("A", "TC-TAB", "Терапия", 0.6),        # дубликат триплета
        ("B", "D-CRM", "Дерматология", 0),       # битая доля
        ("C", "LOZ", "ЛОР", 0.7),
        ("C", "LOZ", "Терапия", 0.6),            # пара C в сумме 130%
    ])
    _, report = parse_directions(path)

    assert report["rows_accepted"] == 2
    assert report["rows_skipped"] == 3
    assert report["rows_read"] == 6
    assert report["rows_read"] == (
        report["rows_accepted"]
        + report["rows_skipped"]
        + report["rows_duplicate"]
    )
    assert report["rows_duplicate"] == 1
    assert report["pairs_over_100"] == 1
    assert report["pairs_count"] == 1          # осталась только пара A


def test_reordered_columns_are_read_by_header_not_position(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    # порядок колонок перепутан
    ws.append(["Направление", "МНН", "ЛФ", "Процент"])
    ws.append(["Дерматология", "CLOBETASOL", "D-CRM", 1.0])
    path = tmp_path / "reordered.xlsx"
    wb.save(path)

    rows, _ = parse_directions(path)

    assert rows == [{
        "mnn": "CLOBETASOL", "lf": "D-CRM",
        "direction": "Дерматология", "share": 1.0,
    }]


def test_extra_leading_column_does_not_shift_fields(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["№", "МНН", "ЛФ", "Направление", "Процент"])
    ws.append([1, "CLOBETASOL", "D-CRM", "Дерматология", 1.0])
    path = tmp_path / "extra.xlsx"
    wb.save(path)

    rows, _ = parse_directions(path)

    assert rows[0]["mnn"] == "CLOBETASOL"
    assert rows[0]["direction"] == "Дерматология"


def test_header_below_a_title_row_is_found(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["Справочник направлений, версия 2026"])
    ws.append([])
    ws.append(["МНН", "ЛФ", "Направление", "Процент"])
    ws.append(["CLOBETASOL", "D-CRM", "Дерматология", 1.0])
    path = tmp_path / "titled.xlsx"
    wb.save(path)

    rows, report = parse_directions(path)

    assert len(rows) == 1
    assert report["rows_read"] == 1       # строка названия не считается данными
    assert report["rows_skipped"] == 0


def test_file_without_header_raises_instead_of_losing_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["CLOBETASOL", "D-CRM", "Дерматология", 1.0])
    ws.append(["ASPIRIN", "TC-TAB", "Кардиология", 1.0])
    path = tmp_path / "noheader.xlsx"
    wb.save(path)

    with pytest.raises(DirectionsParseError):
        parse_directions(path)


@pytest.mark.parametrize("bad", ["nan", "NaN", "inf", "-inf"])
def test_non_finite_share_is_rejected(tmp_path, bad):
    path = make_file(tmp_path, SHEET_NAME, [("A", "TC-TAB", "Терапия", bad)])
    rows, report = parse_directions(path)

    assert rows == []
    assert report["rows_skipped"] == 1


def test_boolean_cell_is_not_treated_as_full_share(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [("A", "TC-TAB", "Терапия", True)])
    rows, report = parse_directions(path)

    assert rows == []
    assert report["rows_skipped"] == 1


def test_comma_decimal_is_accepted(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [("A", "TC-TAB", "Терапия", "0,5")])
    rows, _ = parse_directions(path)

    assert rows[0]["share"] == 0.5


def test_percent_sign_text_is_rejected_not_guessed(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [("A", "TC-TAB", "Терапия", "35%")])
    rows, report = parse_directions(path)

    assert rows == []
    assert report["rows_skipped"] == 1


def test_error_message_lists_found_sheets(tmp_path):
    wb = Workbook()
    wb.active.title = "Первый"
    wb.active.append(["Колонка"])
    wb.create_sheet("Второй").append(["Другая"])
    path = tmp_path / "bad.xlsx"
    wb.save(path)

    with pytest.raises(DirectionsParseError) as exc:
        parse_directions(path)

    text = str(exc.value)
    assert "Первый" in text and "Второй" in text


def test_sheet_with_only_some_required_columns_is_rejected(tmp_path):
    # три колонки из четырёх: мутация «все -> любая» приняла бы такой лист
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(["МНН", "ЛФ", "Направление"])
    ws.append(["A", "TC-TAB", "Терапия"])
    path = tmp_path / "partial.xlsx"
    wb.save(path)

    with pytest.raises(DirectionsParseError):
        parse_directions(path)


def test_empty_lf_is_skipped(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "", "Терапия", 1.0),
        ("B", "TC-TAB", "Терапия", 1.0),
    ])
    rows, report = parse_directions(path)

    assert [r["mnn"] for r in rows] == ["B"]
    assert report["rows_skipped"] == 1


def test_negative_and_non_numeric_shares_are_skipped(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", -0.5),
        ("B", "TC-TAB", "Терапия", "много"),
        ("C", "TC-TAB", "Терапия", 1.0),
    ])
    rows, report = parse_directions(path)

    assert [r["mnn"] for r in rows] == ["C"]
    assert report["rows_skipped"] == 2


def test_tolerance_below_hundred_is_not_incomplete(tmp_path):
    # 0.6 + 0.3 + 0.1 == 0.9999999999999999 в float (не ровно 1.0)
    assert 0.6 + 0.3 + 0.1 != 1.0
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.6),
        ("A", "TC-TAB", "Педиатрия", 0.3),
        ("A", "TC-TAB", "ЛОР", 0.1),
    ])
    _, report = parse_directions(path)

    assert report["pairs_incomplete"] == 0
    assert report["pairs_over_100"] == 0
    assert report["pairs_count"] == 1


def test_tolerance_above_hundred_does_not_drop_pair(tmp_path):
    # 0.05 + 0.55 + 0.3 + 0.1 == 1.0000000000000002 в float (чуть выше 1.0)
    assert 0.05 + 0.55 + 0.3 + 0.1 > 1.0
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.05),
        ("A", "TC-TAB", "Педиатрия", 0.55),
        ("A", "TC-TAB", "ЛОР", 0.3),
        ("A", "TC-TAB", "Кардиология", 0.1),
    ])
    rows, report = parse_directions(path)

    assert len(rows) == 4
    assert report["pairs_over_100"] == 0
    assert report["pairs_incomplete"] == 0


def test_duplicate_with_broken_share_keeps_previous_value(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия", 0.5),
        ("A", "TC-TAB", "Терапия", 0),
    ])
    rows, report = parse_directions(path)

    assert [r["share"] for r in rows] == [0.5]
    assert report["rows_skipped"] == 1
    assert report["rows_duplicate"] == 0


def test_short_row_missing_trailing_cells_is_skipped_not_crashing(tmp_path):
    path = make_file(tmp_path, SHEET_NAME, [
        ("A", "TC-TAB", "Терапия"),             # нет ячейки с долей
        ("B", "TC-TAB", "Терапия", 1.0),
    ])
    rows, report = parse_directions(path)

    assert [r["mnn"] for r in rows] == ["B"]
    assert report["rows_skipped"] == 1
