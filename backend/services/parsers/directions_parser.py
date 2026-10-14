import logging
import math
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger(__name__)

SHEET_NAME = "Длинный формат"

HEADER_SCAN_ROWS = 10

_FIELD_BY_HEADER = {
    "МНН": "mnn",
    "ЛФ": "lf",
    "НАПРАВЛЕНИЕ": "direction",
    "ПРОЦЕНТ": "share",
}


class DirectionsParseError(Exception):
    pass


def _header_map(head: list[str]) -> dict[str, int] | None:
    """Карта поле → индекс колонки, или None если шапка не та."""
    idx: dict[str, int] = {}
    for i, name in enumerate(head):
        field = _FIELD_BY_HEADER.get(name)
        if field is not None and field not in idx:
            idx[field] = i
    return idx if len(idx) == len(_FIELD_BY_HEADER) else None


def _pick_sheet(wb) -> tuple[str, dict[str, int], int]:
    """Лист, карта колонок и номер строки шапки (с единицы)."""
    # Лист с ожидаемым именем проверяется первым, остальные — по шапке.
    candidates = [SHEET_NAME] if SHEET_NAME in wb.sheetnames else []
    candidates += [n for n in wb.sheetnames if n != SHEET_NAME]

    for name in candidates:
        ws = wb[name]
        rows = ws.iter_rows(
            min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True,
        )
        for row_no, values in enumerate(rows, start=1):
            head = [
                str(v).strip().upper() if v is not None else ""
                for v in values
            ]
            cmap = _header_map(head)
            if cmap is not None:
                return name, cmap, row_no

    raise DirectionsParseError(
        f"Не найден лист с колонками МНН, ЛФ, Направление, Процент. "
        f"Ожидался лист «{SHEET_NAME}». "
        f"Найденные листы: {', '.join(wb.sheetnames)}"
    )


def _share(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
    try:
        share = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(share):
        return None
    if share > 1:
        share /= 100
    if share <= 0 or share > 1:
        return None
    return share


def parse_directions(file_path: Path) -> tuple[list[dict], dict]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet, cmap, header_row = _pick_sheet(wb)
        ws = wb[sheet]

        accepted: dict[tuple[str, str, str], dict] = {}
        rows_read = 0
        rows_skipped = 0
        rows_duplicate = 0

        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rows_read += 1

            def cell(field: str):
                i = cmap[field]
                return values[i] if i < len(values) else None

            mnn = str(cell("mnn") or "").strip().upper()
            lf = str(cell("lf") or "").strip().upper()
            direction = str(cell("direction") or "").strip()
            share = _share(cell("share"))

            if not mnn or not lf or not direction or share is None:
                rows_skipped += 1
                continue

            key = (mnn, lf, direction)
            if key in accepted:
                rows_duplicate += 1
            accepted[key] = {
                "mnn": mnn, "lf": lf,
                "direction": direction, "share": share,
            }
    finally:
        wb.close()

    rows = list(accepted.values())

    pair_shares: dict[tuple[str, str], float] = {}
    for r in rows:
        key = (r["mnn"], r["lf"])
        pair_shares[key] = pair_shares.get(key, 0.0) + r["share"]

    # Пара с суммой долей выше 100% ломает инвариант сходимости в сторону
    # раздувания рынка: недобора у неё нет, а переизбыток попал бы в итог
    # лишними деньгами. Отбрасываем целиком, её USD уйдёт в «Не распределено».
    over = {k for k, v in pair_shares.items() if v > 1.0 + 1e-6}
    if over:
        rows = [r for r in rows if (r["mnn"], r["lf"]) not in over]
        rows_skipped += sum(
            1 for r in accepted.values() if (r["mnn"], r["lf"]) in over
        )
        for k in over:
            pair_shares.pop(k, None)

    report = {
        "sheet": sheet,
        "rows_read": rows_read,
        "rows_accepted": len(rows),
        "rows_skipped": rows_skipped,
        "rows_duplicate": rows_duplicate,
        "pairs_count": len(pair_shares),
        "directions_count": len({r["direction"] for r in rows}),
        "forms_count": len({r["lf"] for r in rows}),
        "pairs_incomplete": sum(
            1 for s in pair_shares.values() if abs(s - 1.0) > 1e-6
        ),
        "pairs_over_100": len(over),
    }
    log.info(
        "Справочник направлений разобран: лист «%s», принято %d строк, "
        "отброшено %d, дубликатов %d, пар %d",
        sheet, report["rows_accepted"], rows_skipped, rows_duplicate,
        report["pairs_count"],
    )
    if over:
        log.warning(
            "Отброшено %d пар с суммой долей выше 100%%", len(over),
        )
    return rows, report
